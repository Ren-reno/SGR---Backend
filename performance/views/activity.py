"""CRUD web de Activity: listar, crear y editar (Fase 6, paso 6.1) y
eliminar (Fase 7, Decisión 26). Incluye filtros en el listado y
exportación a Excel (Fase 9).

Capas de control, todas del lado del servidor:
- Permiso de modelo (`PermissionRequiredMixin`): los mismos que ya usa el
  Admin y asigna `seed_sgr` -- Funcionario/Verificador ven y editan,
  solo Administrador crea y elimina; Delegado no tiene ninguno todavía.
  Sin sesión redirige al login; con sesión pero sin permiso responde 403.
- Scoping por Delegación (`DelegationScopedQuerysetMixin`): en el listado
  filtra las filas; en la edición hace que una actividad ajena responda 404.
- Valores escritos: `ActivityForm` acota el desplegable `employee`.
"""
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.contrib.messages.views import SuccessMessageMixin
from django.http import HttpResponse
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, ListView, UpdateView, View
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from organization.models import Employee
from performance.deletion import SoftDeleteView
from performance.forms import ActivityForm
from performance.models import Activity, CatalogItem, Period
from performance.pagination import SessionPaginationMixin
from performance.scoping import DelegationScopedQuerysetMixin, scope_queryset_for_user


class ActivityListView(
    PermissionRequiredMixin, DelegationScopedQuerysetMixin,
    SessionPaginationMixin, ListView,
):
    model = Activity
    template_name = 'performance/activity_list.html'
    context_object_name = 'activities'
    permission_required = 'performance.view_activity'
    delegation_lookup = 'employee__'
    # `ordering` es obligatorio (ver SessionPaginationMixin); `-pk` desempata
    # actividades del mismo día para que ninguna cambie de página al recargar.
    ordering = ('-date', '-pk')

    def get_queryset(self):
        qs = super().get_queryset().select_related(
            'employee', 'period', 'meta', 'attention',
        )
        # Filtros del listado (Fase 9): todos opcionales, por parámetros GET.
        params = self.request.GET
        if params.get('date_from'):
            qs = qs.filter(date__gte=params['date_from'])
        if params.get('date_to'):
            qs = qs.filter(date__lte=params['date_to'])
        if params.get('employee'):
            qs = qs.filter(employee_id=params['employee'])
        if params.get('period'):
            qs = qs.filter(period_id=params['period'])
        if params.get('activity_type'):
            qs = qs.filter(activity_type=params['activity_type'])
        if params.get('attention'):
            qs = qs.filter(attention_id=params['attention'])
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Los desplegables respetan el mismo scoping por Delegación que el
        # listado -- un Funcionario no debe ver, ni en un filtro, empleados
        # de otra Delegación.
        context['filter_employees'] = scope_queryset_for_user(
            Employee.objects.order_by('name'), self.request.user, '',
        )
        context['filter_periods'] = Period.objects.order_by('-start_date')
        context['filter_attentions'] = CatalogItem.objects.filter(
            category=CatalogItem.CATEGORY_ATTENTION, is_active=True,
        ).order_by('name')
        context['filter_activity_types'] = (
            self.get_scoped_queryset()
            .order_by('activity_type')
            .values_list('activity_type', flat=True)
            .distinct()
        )
        context['filter_values'] = self.request.GET
        return context


class _ActivityFormMixin:
    model = Activity
    form_class = ActivityForm
    template_name = 'performance/activity_form.html'
    success_url = reverse_lazy('performance:activity_list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class ActivityCreateView(
    PermissionRequiredMixin, _ActivityFormMixin, SuccessMessageMixin, CreateView,
):
    permission_required = 'performance.add_activity'
    success_message = 'Actividad registrada.'
    extra_context = {'page_title': 'Nueva actividad'}


class ActivityUpdateView(
    PermissionRequiredMixin, DelegationScopedQuerysetMixin,
    _ActivityFormMixin, SuccessMessageMixin, UpdateView,
):
    permission_required = 'performance.change_activity'
    delegation_lookup = 'employee__'
    success_message = 'Actividad actualizada.'
    extra_context = {'page_title': 'Editar actividad'}


class ActivityDeleteView(SoftDeleteView):
    """Borrado lógico (`deleted_at`). Una actividad con evidencias vivas no
    se elimina: `Activity._before_soft_delete` lanza `ProtectedError` y la
    vista base lo muestra como mensaje."""
    model = Activity
    permission_required = 'performance.delete_activity'
    delegation_lookup = 'employee__'
    success_url = reverse_lazy('performance:activity_list')
    success_message = 'Actividad eliminada.'


class ActivityExportView(PermissionRequiredMixin, DelegationScopedQuerysetMixin, View):
    """Exporta a .xlsx las Activity visibles para el usuario actual (Fase 9,
    paso 9.2). Reutiliza get_scoped_queryset() de DelegationScopedQuerysetMixin
    -- mismo scoping por Delegación y mismo soft-delete que ActivityListView,
    no una consulta aparte sin esos filtros.
    """
    model = Activity
    permission_required = 'performance.view_activity'
    delegation_lookup = 'employee__'

    COLUMNS = [
        ('ID', 'pk'),
        ('Fecha', 'date'),
        ('Funcionario', 'employee'),
        ('Período', 'period'),
        ('Meta', 'meta'),
        ('Tipo de actividad', 'activity_type'),
        ('Servicio', 'service'),
        ('Atención', 'attention'),
        ('Subatención', 'sub_attention'),
        ('Descripción de la solicitud', 'request_description'),
        ('Acción realizada', 'action_taken'),
        ('Contacto', 'contact_name'),
        ('Teléfono', 'contact_phone'),
        ('Estado', 'status'),
    ]

    def get(self, request, *args, **kwargs):
        queryset = self.get_scoped_queryset().select_related(
            'employee', 'period', 'meta', 'attention',
        ).order_by('-date', '-pk')

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Actividades'

        headers = [label for label, _ in self.COLUMNS]
        sheet.append(headers)
        for col_index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col_index)].width = 22

        for activity in queryset.iterator():
            row = []
            for _, field in self.COLUMNS:
                value = activity.pk if field == 'pk' else getattr(activity, field)
                row.append(str(value) if value is not None else '')
            sheet.append(row)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        filename = f'actividades_{timezone.now():%Y%m%d_%H%M}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        workbook.save(response)
        return response